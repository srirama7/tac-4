"""
Unit tests for export_processor module.
Tests CSV, JSON, and Excel export functionality with various data types and edge cases.
"""

import pytest
import json
import csv
from io import StringIO, BytesIO
from openpyxl import load_workbook
from core.export_processor import (
    export_to_csv,
    export_to_json,
    export_to_excel,
    generate_export_filename,
    validate_export_data
)


class TestCSVExport:
    """Tests for CSV export functionality"""

    def test_basic_csv_export(self):
        """Test basic CSV export with simple data"""
        data = [
            {"id": 1, "name": "Alice", "age": 30},
            {"id": 2, "name": "Bob", "age": 25}
        ]
        columns = ["id", "name", "age"]

        result = export_to_csv(data, columns)

        # Decode and parse CSV
        csv_text = result.decode('utf-8')
        reader = csv.DictReader(StringIO(csv_text))
        rows = list(reader)

        assert len(rows) == 2
        assert rows[0]["id"] == "1"
        assert rows[0]["name"] == "Alice"
        assert rows[1]["name"] == "Bob"

    def test_csv_with_special_characters(self):
        """Test CSV export with commas, quotes, and newlines"""
        data = [
            {"id": 1, "description": "Hello, World"},
            {"id": 2, "description": 'Quote: "test"'},
            {"id": 3, "description": "Line\nbreak"}
        ]
        columns = ["id", "description"]

        result = export_to_csv(data, columns)
        csv_text = result.decode('utf-8')

        # Verify CSV is valid
        reader = csv.DictReader(StringIO(csv_text))
        rows = list(reader)

        assert len(rows) == 3
        assert rows[0]["description"] == "Hello, World"
        assert rows[1]["description"] == 'Quote: "test"'

    def test_csv_with_null_values(self):
        """Test CSV export handles null/None values"""
        data = [
            {"id": 1, "name": "Alice", "email": None},
            {"id": 2, "name": None, "email": "bob@example.com"}
        ]
        columns = ["id", "name", "email"]

        result = export_to_csv(data, columns)
        csv_text = result.decode('utf-8')

        reader = csv.DictReader(StringIO(csv_text))
        rows = list(reader)

        assert len(rows) == 2
        assert rows[0]["email"] == ""
        assert rows[1]["name"] == ""

    def test_csv_with_empty_data(self):
        """Test CSV export with empty result set"""
        data = []
        columns = ["id", "name"]

        result = export_to_csv(data, columns)
        csv_text = result.decode('utf-8')

        # Should have header only
        lines = csv_text.strip().split('\n')
        assert len(lines) == 1
        assert "id" in lines[0] and "name" in lines[0]

    def test_csv_with_numeric_types(self):
        """Test CSV export with integers and floats"""
        data = [
            {"id": 1, "price": 19.99, "quantity": 5},
            {"id": 2, "price": 29.95, "quantity": 10}
        ]
        columns = ["id", "price", "quantity"]

        result = export_to_csv(data, columns)
        csv_text = result.decode('utf-8')

        reader = csv.DictReader(StringIO(csv_text))
        rows = list(reader)

        assert len(rows) == 2
        assert rows[0]["price"] == "19.99"
        assert rows[1]["quantity"] == "10"


class TestJSONExport:
    """Tests for JSON export functionality"""

    def test_basic_json_export(self):
        """Test basic JSON export"""
        data = [
            {"id": 1, "name": "Alice", "active": True},
            {"id": 2, "name": "Bob", "active": False}
        ]
        columns = ["id", "name", "active"]

        result = export_to_json(data, columns)

        # Parse JSON
        parsed = json.loads(result.decode('utf-8'))

        assert len(parsed) == 2
        assert parsed[0]["id"] == 1
        assert parsed[0]["name"] == "Alice"
        assert parsed[0]["active"] is True

    def test_json_with_nested_structures(self):
        """Test JSON export with nested data"""
        data = [
            {"id": 1, "meta": {"views": 100, "likes": 50}},
            {"id": 2, "meta": {"views": 200, "likes": 75}}
        ]
        columns = ["id", "meta"]

        result = export_to_json(data, columns)
        parsed = json.loads(result.decode('utf-8'))

        assert parsed[0]["meta"]["views"] == 100
        assert parsed[1]["meta"]["likes"] == 75

    def test_json_with_special_characters(self):
        """Test JSON export with unicode and special characters"""
        data = [
            {"id": 1, "text": "Hello 世界"},
            {"id": 2, "text": "Émoji: 😀"}
        ]
        columns = ["id", "text"]

        result = export_to_json(data, columns)
        parsed = json.loads(result.decode('utf-8'))

        assert parsed[0]["text"] == "Hello 世界"
        assert parsed[1]["text"] == "Émoji: 😀"

    def test_json_with_null_values(self):
        """Test JSON export preserves null values"""
        data = [
            {"id": 1, "name": "Alice", "email": None},
            {"id": 2, "name": None, "email": "bob@example.com"}
        ]
        columns = ["id", "name", "email"]

        result = export_to_json(data, columns)
        parsed = json.loads(result.decode('utf-8'))

        assert parsed[0]["email"] is None
        assert parsed[1]["name"] is None

    def test_json_with_empty_data(self):
        """Test JSON export with empty result set"""
        data = []
        columns = ["id", "name"]

        result = export_to_json(data, columns)
        parsed = json.loads(result.decode('utf-8'))

        assert parsed == []
        assert isinstance(parsed, list)


class TestExcelExport:
    """Tests for Excel export functionality"""

    def test_basic_excel_export(self):
        """Test basic Excel export"""
        data = [
            {"id": 1, "name": "Alice", "age": 30},
            {"id": 2, "name": "Bob", "age": 25}
        ]
        columns = ["id", "name", "age"]

        result = export_to_excel(data, columns)

        # Load workbook from bytes
        wb = load_workbook(BytesIO(result))
        ws = wb.active

        # Check headers
        assert ws.cell(1, 1).value == "id"
        assert ws.cell(1, 2).value == "name"
        assert ws.cell(1, 3).value == "age"

        # Check data
        assert ws.cell(2, 1).value == 1
        assert ws.cell(2, 2).value == "Alice"
        assert ws.cell(2, 3).value == 30
        assert ws.cell(3, 2).value == "Bob"

    def test_excel_with_header_formatting(self):
        """Test Excel export includes header formatting"""
        data = [{"id": 1, "name": "Test"}]
        columns = ["id", "name"]

        result = export_to_excel(data, columns)

        wb = load_workbook(BytesIO(result))
        ws = wb.active

        # Check header has bold font
        header_cell = ws.cell(1, 1)
        assert header_cell.font.bold is True

    def test_excel_with_null_values(self):
        """Test Excel export handles null values"""
        data = [
            {"id": 1, "name": "Alice", "email": None},
            {"id": 2, "name": None, "email": "bob@example.com"}
        ]
        columns = ["id", "name", "email"]

        result = export_to_excel(data, columns)

        wb = load_workbook(BytesIO(result))
        ws = wb.active

        # In Excel, we convert None to empty string for better display
        # When checking, empty cells have None value but display as empty
        assert ws.cell(2, 3).value in ("", None)  # None becomes empty string
        assert ws.cell(3, 2).value in ("", None)

    def test_excel_with_empty_data(self):
        """Test Excel export with empty result set"""
        data = []
        columns = ["id", "name"]

        result = export_to_excel(data, columns)

        wb = load_workbook(BytesIO(result))
        ws = wb.active

        # Should have headers only
        assert ws.cell(1, 1).value == "id"
        assert ws.cell(1, 2).value == "name"
        assert ws.cell(2, 1).value is None  # No data rows

    def test_excel_with_numeric_types(self):
        """Test Excel export with numeric data types"""
        data = [
            {"id": 1, "price": 19.99, "quantity": 5},
            {"id": 2, "price": 29.95, "quantity": 10}
        ]
        columns = ["id", "price", "quantity"]

        result = export_to_excel(data, columns)

        wb = load_workbook(BytesIO(result))
        ws = wb.active

        assert ws.cell(2, 2).value == 19.99
        assert ws.cell(3, 3).value == 10


class TestFilenameGeneration:
    """Tests for filename generation"""

    def test_basic_filename_generation(self):
        """Test basic filename with timestamp"""
        filename = generate_export_filename("csv")

        assert filename.startswith("query_results_")
        assert filename.endswith(".csv")
        assert len(filename) > 15  # Should include timestamp

    def test_filename_with_table_name(self):
        """Test filename includes table name"""
        filename = generate_export_filename("csv", "users")

        assert filename.startswith("users_")
        assert filename.endswith(".csv")

    def test_filename_format_mapping(self):
        """Test correct file extensions for each format"""
        csv_file = generate_export_filename("csv")
        json_file = generate_export_filename("json")
        excel_file = generate_export_filename("excel")

        assert csv_file.endswith(".csv")
        assert json_file.endswith(".json")
        assert excel_file.endswith(".xlsx")

    def test_filename_sanitization(self):
        """Test filename sanitizes special characters"""
        filename = generate_export_filename("csv", "user$_data!@#")

        # Should remove special characters
        assert "$" not in filename
        assert "!" not in filename
        assert "@" not in filename


class TestDataValidation:
    """Tests for data validation"""

    def test_valid_data_passes_validation(self):
        """Test validation passes for valid data"""
        data = [
            {"id": 1, "name": "Alice"},
            {"id": 2, "name": "Bob"}
        ]
        columns = ["id", "name"]

        # Should not raise exception
        validate_export_data(data, columns)

    def test_empty_data_passes_validation(self):
        """Test validation allows empty data"""
        data = []
        columns = ["id", "name"]

        # Should not raise exception
        validate_export_data(data, columns)

    def test_invalid_data_type_raises_error(self):
        """Test validation fails for non-list data"""
        with pytest.raises(ValueError, match="Data must be a list"):
            validate_export_data("not a list", ["id"])

    def test_invalid_columns_type_raises_error(self):
        """Test validation fails for non-list columns"""
        with pytest.raises(ValueError, match="Columns must be a list"):
            validate_export_data([], "not a list")

    def test_non_dict_rows_raise_error(self):
        """Test validation fails for non-dictionary rows"""
        with pytest.raises(ValueError, match="Data rows must be dictionaries"):
            validate_export_data([1, 2, 3], ["id"])


class TestEdgeCases:
    """Tests for edge cases and error handling"""

    def test_large_dataset_csv(self):
        """Test CSV export with large dataset"""
        data = [{"id": i, "value": f"item_{i}"} for i in range(1000)]
        columns = ["id", "value"]

        result = export_to_csv(data, columns)

        csv_text = result.decode('utf-8')
        reader = csv.DictReader(StringIO(csv_text))
        rows = list(reader)

        assert len(rows) == 1000

    def test_unicode_in_all_formats(self):
        """Test unicode handling in all export formats"""
        data = [{"text": "Hello 世界 🌍"}]
        columns = ["text"]

        # CSV
        csv_result = export_to_csv(data, columns)
        assert "世界" in csv_result.decode('utf-8')

        # JSON
        json_result = export_to_json(data, columns)
        parsed = json.loads(json_result.decode('utf-8'))
        assert parsed[0]["text"] == "Hello 世界 🌍"

        # Excel
        excel_result = export_to_excel(data, columns)
        wb = load_workbook(BytesIO(excel_result))
        ws = wb.active
        assert "世界" in str(ws.cell(2, 1).value)

    def test_very_long_text_fields(self):
        """Test export with very long text fields"""
        long_text = "A" * 10000
        data = [{"id": 1, "text": long_text}]
        columns = ["id", "text"]

        # Should not crash
        csv_result = export_to_csv(data, columns)
        assert len(csv_result) > 10000

        json_result = export_to_json(data, columns)
        assert len(json_result) > 10000

        excel_result = export_to_excel(data, columns)
        assert len(excel_result) > 0
