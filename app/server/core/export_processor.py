"""
Export processor module for converting query results to various formats.
Supports CSV, JSON, and Excel (XLSX) export formats.
"""

import csv
import json
import logging
from io import BytesIO, StringIO
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from datetime import datetime

logger = logging.getLogger(__name__)


def export_to_csv(data: List[Dict[str, Any]], columns: List[str]) -> bytes:
    """
    Export data to CSV format.

    Args:
        data: List of dictionaries containing the data to export
        columns: List of column names

    Returns:
        bytes: CSV file content as bytes

    Raises:
        ValueError: If data structure is invalid
    """
    try:
        # Create string buffer for CSV
        output = StringIO()

        # Write CSV with proper escaping
        writer = csv.DictWriter(
            output,
            fieldnames=columns,
            quoting=csv.QUOTE_MINIMAL,
            lineterminator='\n'
        )

        # Write header
        writer.writeheader()

        # Write data rows
        for row in data:
            # Handle None values
            cleaned_row = {k: (v if v is not None else '') for k, v in row.items()}
            writer.writerow(cleaned_row)

        # Convert to bytes
        csv_content = output.getvalue()
        output.close()

        logger.info(f"Exported {len(data)} rows to CSV format")
        return csv_content.encode('utf-8')

    except Exception as e:
        logger.error(f"CSV export error: {str(e)}")
        raise ValueError(f"Failed to export to CSV: {str(e)}")


def export_to_json(data: List[Dict[str, Any]], columns: List[str]) -> bytes:
    """
    Export data to JSON format.

    Args:
        data: List of dictionaries containing the data to export
        columns: List of column names (unused, for consistency)

    Returns:
        bytes: JSON file content as bytes

    Raises:
        ValueError: If data structure is invalid or cannot be serialized
    """
    try:
        # Create JSON with proper formatting
        json_content = json.dumps(
            data,
            indent=2,
            ensure_ascii=False,
            default=str  # Handle datetime and other non-serializable types
        )

        logger.info(f"Exported {len(data)} rows to JSON format")
        return json_content.encode('utf-8')

    except TypeError as e:
        logger.error(f"JSON serialization error: {str(e)}")
        raise ValueError(f"Failed to serialize data to JSON: {str(e)}")
    except Exception as e:
        logger.error(f"JSON export error: {str(e)}")
        raise ValueError(f"Failed to export to JSON: {str(e)}")


def export_to_excel(data: List[Dict[str, Any]], columns: List[str]) -> bytes:
    """
    Export data to Excel (XLSX) format using openpyxl.

    Args:
        data: List of dictionaries containing the data to export
        columns: List of column names

    Returns:
        bytes: Excel file content as bytes

    Raises:
        ValueError: If data structure is invalid or Excel generation fails
    """
    try:
        # Create workbook and active sheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Query Results"

        # Style for header row
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

        # Write headers
        for col_idx, column_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = column_name
            cell.font = header_font
            cell.fill = header_fill

        # Write data rows
        for row_idx, row_data in enumerate(data, start=2):
            for col_idx, column_name in enumerate(columns, start=1):
                value = row_data.get(column_name)

                # Handle None values
                if value is None:
                    value = ""
                # Convert datetime objects to strings
                elif isinstance(value, datetime):
                    value = value.isoformat()

                ws.cell(row=row_idx, column=col_idx, value=value)

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                try:
                    cell_value = str(cell.value) if cell.value is not None else ""
                    if len(cell_value) > max_length:
                        max_length = len(cell_value)
                except:
                    pass

            # Set width with some padding, max 50
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Save to BytesIO buffer
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        excel_content = output.read()
        output.close()

        logger.info(f"Exported {len(data)} rows to Excel format")
        return excel_content

    except Exception as e:
        logger.error(f"Excel export error: {str(e)}")
        raise ValueError(f"Failed to export to Excel: {str(e)}")


def generate_export_filename(format: str, table_name: str = None) -> str:
    """
    Generate a filename for export with timestamp.

    Args:
        format: Export format ('csv', 'json', 'excel')
        table_name: Optional table name to include in filename

    Returns:
        str: Generated filename
    """
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    # Map format to extension
    extensions = {
        'csv': 'csv',
        'json': 'json',
        'excel': 'xlsx'
    }

    ext = extensions.get(format, format)

    if table_name:
        # Sanitize table name
        safe_table_name = "".join(c for c in table_name if c.isalnum() or c in ('_', '-'))
        return f"{safe_table_name}_{timestamp}.{ext}"

    return f"query_results_{timestamp}.{ext}"


def validate_export_data(data: List[Dict[str, Any]], columns: List[str]) -> None:
    """
    Validate export data structure.

    Args:
        data: List of dictionaries to validate
        columns: Expected column names

    Raises:
        ValueError: If data structure is invalid
    """
    if not isinstance(data, list):
        raise ValueError("Data must be a list")

    if not isinstance(columns, list):
        raise ValueError("Columns must be a list")

    if len(data) > 0 and not isinstance(data[0], dict):
        raise ValueError("Data rows must be dictionaries")

    # Validate column consistency (optional - handle missing columns gracefully)
    if len(data) > 0:
        first_row_keys = set(data[0].keys())
        if not set(columns).issubset(first_row_keys | {None}):
            logger.warning(f"Column mismatch detected. Expected: {columns}, Got: {list(first_row_keys)}")
